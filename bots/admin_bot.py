# bots/admin_bot.py (TOPs fix + Excel sheet rename)
import asyncio
from collections import defaultdict
from datetime import datetime, date
from pathlib import Path

from aiogram import Bot, Dispatcher, Router, F
from aiogram.filters import CommandStart, Command, StateFilter
from aiogram.types import Message, CallbackQuery, FSInputFile
from aiogram.fsm.state import StatesGroup, State
from aiogram.fsm.context import FSMContext
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.utils.keyboard import InlineKeyboardBuilder
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from common.config import get_admin_token, get_admin_access_pass, get_doctor_token
from common.logging_config import setup_logger
from common.db import (
    migrate,
    admin_is_authorized,
    admin_authorize,

    list_doctors_with_counts,
    search_doctors_prefix,

    list_all_patients,
    list_doctor_patients_visit_status,

    list_unvisited_all,
    list_unvisited_by_doctor,
    mark_referrals_visited,

    list_ready_to_settle_all,
    list_ready_to_settle_by_doctor,
    settle_referrals,
    get_referrals_details,

    export_doctors_overview,
    export_all_referrals,

    list_patients_with_ref_counts,
    delete_patients,
    list_patients_by_doctor_distinct,
    search_patients_prefix,
    list_referrals_for_patient,

    # TOPs
    top_doctors_by_referrals,
    top_doctors_by_visits,
)

ROOT = Path(__file__).resolve().parents[1]
log = setup_logger("admin_bot", ROOT / "logs" / "admin.log")

router = Router()
doctor_notify_bot: Bot | None = None


# -------- utils --------

def _fmt_help() -> str:
    return (
        "Админ-бот. Доступные действия:\n"
        "• 👨‍⚕️ Врачи — список врачей и их направлений\n"
        "• 🔎 Поиск врача — быстрый поиск по ФИО\n"
        "• 🔎 Поиск пациента — быстрый поиск по ФИО\n"
        "• 👥 Все пациенты — полный список пациентов клиники\n"
        "• 📋 Списки пациентов — вывод за всё время/месяц/год+месяц\n"
        "• 🏆 Рейтинг врачей — ТОП по направлениям/визитам\n"
        "• 🕓 Отметить визит — кого ещё не отметили\n"
        "• 💳 Расчёт за визит — все отметившиеся или по врачу\n"
        "• ✅ Рассчитанные — за текущий месяц\n"
        "• 🗑 Удалить пациента(ов) — по номерам из списка\n"
        "• 📤 Экспорт — Excel-файл со сводкой, направлениями и рейтингом\n"
    )


def _kb_main() -> InlineKeyboardBuilder:
    kb = InlineKeyboardBuilder()
    kb.button(text="👨‍⚕️ Врачи", callback_data="adm:doctors")
    kb.button(text="🔎 Поиск врача", callback_data="adm:find")
    kb.button(text="🔎 Поиск пациента", callback_data="adm:find_patient")
    kb.button(text="👥 Все пациенты", callback_data="adm:patients_all")
    kb.button(text="📋 Списки пациентов", callback_data="adm:plist_menu")
    kb.button(text="🏆 Рейтинг врачей", callback_data="adm:tops_menu")
    kb.button(text="🕓 Отметить визит", callback_data="adm:visits_menu")
    kb.button(text="💳 Расчёт за визит", callback_data="adm:settle_menu")
    kb.button(text="✅ Рассчитанные", callback_data="adm:settled_current")
    kb.button(text="📤 Экспорт (Excel)", callback_data="adm:export")
    kb.adjust(1)
    return kb


def _kb_help() -> InlineKeyboardBuilder:
    kb = _kb_main()
    kb.button(text="🗑 Удалить пациента(ов)", callback_data="adm:del_patients")
    kb.adjust(1)
    return kb


def _kb_back_menu() -> InlineKeyboardBuilder:
    kb = InlineKeyboardBuilder()
    kb.button(text="⬅️ В меню", callback_data="adm:back")
    kb.adjust(1)
    return kb


async def _send_chunked(msg: Message, title: str, lines: list[str]):
    if not lines:
        await msg.answer(f"{title}\n(пусто)")
        return
    chunk = ""
    for line in lines:
        if len(chunk) + len(line) + 1 > 3500:
            await msg.answer(f"{title}\n{chunk.strip()}")
            chunk = ""
        chunk += line + "\n"
    if chunk:
        await msg.answer(f"{title}\n{chunk.strip()}")


def _fmt_ru_date(iso_ts: str | None) -> str:
    if not iso_ts:
        return "—"
    try:
        return datetime.fromisoformat(iso_ts).strftime("%d.%m.%Y")
    except Exception:
        return "—"


def _fmt_birth_ru(birth_iso: str | None) -> str:
    if not birth_iso:
        return "— дата не указана"
    try:
        return datetime.strptime(birth_iso, "%Y-%m-%d").strftime("%d.%m.%Y")
    except Exception:
        return "— дата не указана"


def _expand_numbers(spec: str, max_n: int) -> list[int]:
    result: set[int] = set()
    s = (spec or "").replace(";", ",").replace("–", "-").replace("—", "-")
    tokens = [t.strip() for t in s.split(",") if t.strip()]
    for t in tokens:
        if "-" in t:
            a, b = t.split("-", 1)
            if not a.isdigit() or not b.isdigit():
                continue
            x, y = int(a), int(b)
            if x > y:
                x, y = y, x
            for k in range(x, y + 1):
                if 1 <= k <= max_n:
                    result.add(k)
        else:
            if t.isdigit():
                k = int(t)
                if 1 <= k <= max_n:
                    result.add(k)
    return sorted(result)


def _autosize(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            v = cell.value
            l = len(str(v)) if v is not None else 0
            if l > max_len:
                max_len = l
        ws.column_dimensions[col_letter].width = min(max_len + 2, 60)


# -------- FSM --------

class DelDoctor(StatesGroup):
    waiting_number = State()
    confirm = State()


class FindDoctor(StatesGroup):
    waiting_query = State()


class MarkVisited(StatesGroup):
    waiting_numbers = State()
    confirm = State()


class MarkSettled(StatesGroup):
    waiting_numbers = State()
    confirm = State()


class DelPatients(StatesGroup):
    waiting_numbers = State()
    confirm = State()


class ViewDocPatients(StatesGroup):
    waiting_number = State()


class FindPatient(StatesGroup):
    waiting_query = State()


class WhoDirected(StatesGroup):
    waiting_number = State()


class PListChoose(StatesGroup):
    waiting_year = State()
    waiting_month = State()
    waiting_number = State()


class TopChoose(StatesGroup):
    waiting_year = State()
    waiting_month = State()


# -------- auth/help --------

@router.message(CommandStart())
async def start_handler(msg: Message):
    if admin_is_authorized(msg.chat.id):
        await msg.answer("✅ Авторизация успешна.\nℹ️ Полное меню — команда /help")
        await msg.answer("Выберите действие:", reply_markup=_kb_main().as_markup())
    else:
        await msg.answer("Отправьте секретный пароль для доступа.")


@router.message(Command("help"))
async def help_handler(msg: Message):
    if not admin_is_authorized(msg.chat.id):
        await msg.answer("Сначала авторизуйтесь — отправьте секретный пароль.")
        return
    await msg.answer(_fmt_help(), reply_markup=_kb_help().as_markup())


@router.message(StateFilter(None))
async def auth_or_ignore(msg: Message):
    if admin_is_authorized(msg.chat.id):
        return
    secret = get_admin_access_pass()
    if msg.text and msg.text.strip() == secret:
        admin_authorize(msg.chat.id)
        await msg.answer("✅ Авторизация успешна.\nℹ️ Полное меню — команда /help")
        await msg.answer("Выберите действие:", reply_markup=_kb_main().as_markup())
        log.info(f"Admin chat authorized: {msg.chat.id}")


# -------- Врачи / поиск / пациенты врача --------

@router.callback_query(F.data == "adm:doctors")
async def list_doctors(cb: CallbackQuery, state: FSMContext):
    data = list_doctors_with_counts()
    if not data:
        await cb.message.answer("Список врачей пуст.")
        await cb.answer()
        return

    lines = [f"{i}. {d['full_name']} — направлений: {d['referrals_count']}" for i, d in enumerate(data, 1)]
    await cb.message.answer("👨‍⚕️ Список врачей:")
    await _send_chunked(cb.message, "Врачи:", lines)
    await state.update_data(last_doctor_list=[int(d["doctor_id"]) for d in data], last_doctor_lines=lines)

    kb = InlineKeyboardBuilder()
    kb.button(text="👁 Пациенты врача", callback_data="adm:doc_patients")
    kb.adjust(1)
    await cb.message.answer("Действия со списком:", reply_markup=kb.as_markup())
    await cb.answer()


@router.callback_query(F.data == "adm:doc_patients")
async def doc_patients_start(cb: CallbackQuery, state: FSMContext):
    await cb.message.answer("Введите номер врача из последнего списка:", reply_markup=_kb_back_menu().as_markup())
    await state.set_state(ViewDocPatients.waiting_number)
    await cb.answer()


@router.message(ViewDocPatients.waiting_number)
async def doc_patients_show(msg: Message, state: FSMContext):
    data = await state.get_data()
    last = data.get("last_doctor_list") or []
    text = (msg.text or "").strip()
    if not text.isdigit():
        await msg.answer("Ожидался номер врача. Попробуйте ещё.", reply_markup=_kb_back_menu().as_markup())
        return
    idx = int(text)
    if idx < 1 or idx > len(last):
        await msg.answer("Неверный номер врача.", reply_markup=_kb_back_menu().as_markup())
        return
    doctor_id = last[idx - 1]
    try:
        rows = list_doctor_patients_visit_status(doctor_id)
    except Exception as e:
        await msg.answer("Не удалось получить список пациентов врача (ошибка).", reply_markup=_kb_back_menu().as_markup())
        try:
            log.exception(f"doc_patients_show failed for doctor_id={doctor_id}: {e}")
        except Exception:
            pass
        return

    lines = []
    for i, p in enumerate(rows or [], 1):
        v = "был" if int(p.get("visited_any") or 0) == 1 else "не был"
        vdate = _fmt_ru_date(p.get("last_visited_at"))
        lines.append(f"{i}. {p['full_name']} — {_fmt_birth_ru(p.get('birth_date'))} — визит: {v} ({vdate})")
    await _send_chunked(msg, "Пациенты врача:", lines or ["(нет)"])
    await state.clear()
    await msg.answer("Выберите действие:", reply_markup=_kb_main().as_markup())


@router.callback_query(F.data == "adm:find")
async def find_start(cb: CallbackQuery, state: FSMContext):
    await cb.message.answer("Введите начало ФИО врача (например, «Ива»).", reply_markup=_kb_back_menu().as_markup())
    await state.set_state(FindDoctor.waiting_query)
    await cb.answer()


@router.message(FindDoctor.waiting_query)
async def find_apply(msg: Message, state: FSMContext):
    q = (msg.text or "").strip()
    res = search_doctors_prefix(q)
    if not res:
        await msg.answer("Ничего не найдено.")
        await msg.answer("Выберите действие:", reply_markup=_kb_main().as_markup())
        await state.clear()
        return

    lines = [f"{i}. {d['full_name']} — направлений: {d['referrals_count']}" for i, d in enumerate(res, 1)]
    await _send_chunked(msg, f"Результат поиска «{q}»: ", lines)
    await state.update_data(last_doctor_list=[int(d["doctor_id"]) for d in res], last_doctor_lines=lines)
    await msg.answer("Выберите действие:", reply_markup=_kb_main().as_markup())
    await state.clear()


# -------- Все пациенты + кто направлял --------

@router.callback_query(F.data == "adm:patients_all")
async def patients_all(cb: CallbackQuery, state: FSMContext):
    rows = list_all_patients()
    lines = [f"{i}. {p['full_name']} — {_fmt_birth_ru(p.get('birth_date'))}" for i, p in enumerate(rows, 1)]
    await _send_chunked(cb.message, "Все пациенты:", lines)

    id_map = [int(p["id"]) for p in rows]
    await state.update_data(all_patients_map=id_map)

    await cb.message.answer(
        "Чтобы посмотреть, кто направлял пациента — отправьте номер из списка.\n"
        "Или нажмите «⬅️ В меню».",
        reply_markup=_kb_back_menu().as_markup()
    )
    await state.set_state(WhoDirected.waiting_number)
    await cb.answer()


@router.message(WhoDirected.waiting_number)
async def who_directed(msg: Message, state: FSMContext):
    text = (msg.text or "").strip().lower()
    if text in {"отмена", "cancel", "stop"}:
        await state.clear()
        await msg.answer("Отменено.")
        await msg.answer("Выберите действие:", reply_markup=_kb_main().as_markup())
        return

    data = await state.get_data()
    ids: list[int] = data.get("all_patients_map") or []
    if not text.isdigit():
        await msg.answer("Ожидался номер пациента из списка.", reply_markup=_kb_back_menu().as_markup())
        return
    idx = int(text)
    if idx < 1 or idx > len(ids):
        await msg.answer("Неверный номер пациента.", reply_markup=_kb_back_menu().as_markup())
        return

    patient_id = ids[idx - 1]
    refs = list_referrals_for_patient(patient_id)
    if not refs:
        await msg.answer("Для пациента направлений пока нет.")
        await state.clear()
        await msg.answer("Выберите действие:", reply_markup=_kb_main().as_markup())
        return

    lines = []
    for r in refs:
        doc = r["doctor_full_name"]
        dadd = _fmt_ru_date(r.get("created_at"))
        visited = "был" if int(r.get("visited") or 0) == 1 else "не был"
        settled = "да" if int(r.get("settled") or 0) == 1 else "—"
        vdate = _fmt_ru_date(r.get("visited_at"))
        sdate = _fmt_ru_date(r.get("settled_at"))
        lines.append(f"• {doc} — направление: {dadd} — визит: {visited} ({vdate}) — расчёт: {settled} ({sdate})")
    await _send_chunked(msg, "Кто направлял:", lines)
    await state.clear()
    await msg.answer("Выберите действие:", reply_markup=_kb_main().as_markup())


# -------- Рейтинг врачей (ТОПы) --------

@router.callback_query(F.data == "adm:tops_menu")
async def tops_menu(cb: CallbackQuery, state: FSMContext):
    kb = InlineKeyboardBuilder()
    kb.button(text="ТОП по направлениям", callback_data="adm:top:ref")
    kb.button(text="ТОП по визитам", callback_data="adm:top:vis")
    kb.button(text="⬅️ Назад", callback_data="adm:back")
    kb.adjust(1)
    await cb.message.answer("Что показать?", reply_markup=kb.as_markup())
    await cb.answer()


@router.callback_query(F.data.in_({"adm:top:ref", "adm:top:vis"}))
async def top_period_menu(cb: CallbackQuery, state: FSMContext):
    metric = cb.data.split(":")[-1]  # ref / vis
    await state.update_data(top_metric=metric)
    kb = InlineKeyboardBuilder()
    kb.button(text="Весь период", callback_data="adm:top:p:all")
    kb.button(text="Текущий месяц", callback_data="adm:top:p:curr")
    kb.button(text="Выбрать год и месяц", callback_data="adm:top:p:choose")
    kb.button(text="⬅️ Назад", callback_data="adm:tops_menu")
    kb.adjust(1)
    title = "ТОП по направлениям" if metric == "ref" else "ТОП по визитам"
    await cb.message.answer(f"{title}: выберите период", reply_markup=kb.as_markup())
    await cb.answer()


def _lines_for_top(rows: list[dict]) -> list[str]:
    lines = []
    for i, r in enumerate(rows, 1):
        lines.append(f"{i}. {r['doctor_full_name']} — {int(r['cnt'] or 0)}")
    return lines


async def _show_top(cb: CallbackQuery, state: FSMContext, year: int | None, month: int | None):
    data = await state.get_data()
    metric = data.get("top_metric")
    if metric == "ref":
        rows = top_doctors_by_referrals(year, month, limit=100)
        title = "ТОП по направлениям"
    else:
        rows = top_doctors_by_visits(year, month, limit=100)
        title = "ТОП по визитам"

    if year and month:
        title += f" — {month:02d}.{year}"
    elif year:
        title += f" — {year}"
    else:
        title += " — весь период"

    await _send_chunked(cb.message, title + ":", _lines_for_top(rows))


@router.callback_query(F.data == "adm:top:p:all")
async def top_all(cb: CallbackQuery, state: FSMContext):
    await _show_top(cb, state, None, None)
    await cb.answer()


@router.callback_query(F.data == "adm:top:p:curr")
async def top_curr(cb: CallbackQuery, state: FSMContext):
    t = date.today()
    await _show_top(cb, state, t.year, t.month)
    await cb.answer()


@router.callback_query(F.data == "adm:top:p:choose")
async def top_choose_year(cb: CallbackQuery, state: FSMContext):
    rows = export_all_referrals()
    years = sorted({int((r.get("created_at") or "")[:4]) for r in rows if (r.get("created_at") or "")[:4].isdigit()})
    if not years:
        years = [date.today().year]
    kb = InlineKeyboardBuilder()
    for y in years:
        kb.button(text=str(y), callback_data=f"adm:top:y:{y}")
    kb.button(text="⬅️ Назад", callback_data="adm:tops_menu")
    kb.adjust(3)
    await cb.message.answer("Выберите год:", reply_markup=kb.as_markup())
    await cb.answer()


@router.callback_query(F.data.startswith("adm:top:y:"))
async def top_choose_month(cb: CallbackQuery, state: FSMContext):
    year = int(cb.data.split(":")[-1])
    await state.update_data(top_year=year)
    kb = InlineKeyboardBuilder()
    for m in range(1, 13):
        kb.button(text=f"{m:02d}", callback_data=f"adm:top:m:{m:02d}")
    kb.button(text="⬅️ Назад", callback_data="adm:top:p:choose")
    kb.adjust(6)
    await cb.message.answer(f"Выбран год {year}. Выберите месяц:", reply_markup=kb.as_markup())
    await cb.answer()


@router.callback_query(F.data.startswith("adm:top:m:"))
async def top_show_month(cb: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    year = int(data.get("top_year"))
    month = int(cb.data.split(":")[-1])
    await _show_top(cb, state, year, month)
    await cb.answer()


# -------- Отметить визит --------

@router.callback_query(F.data == "adm:visits_menu")
async def visits_menu(cb: CallbackQuery, state: FSMContext):
    kb = InlineKeyboardBuilder()
    kb.button(text="Все не отмеченные визиты", callback_data="adm:vis_all")
    kb.button(text="Не отмеченные — по врачу", callback_data="adm:vis_choose_doctor")
    kb.button(text="⬅️ Назад", callback_data="adm:back")
    kb.adjust(1)
    await cb.message.answer("Выберите режим отметки визита:", reply_markup=kb.as_markup())
    await cb.answer()


def _lines_for_unvisited(items: list[dict]) -> list[str]:
    lines = []
    for i, r in enumerate(items, 1):
        fio = r["patient_full_name"]
        bd = _fmt_birth_ru(r.get("patient_birth_date"))
        dadd = _fmt_ru_date(r.get("created_at"))
        doc = r["doctor_full_name"]
        lines.append(f"{i}. {fio} {bd} — направил: {doc} — дата: {dadd}")
    return lines


async def _start_mark_visit_flow(msg: Message, items: list[dict], state: FSMContext, title: str):
    if not items:
        await msg.answer(f"{title}\nСписок пуст.")
        await msg.answer("Выберите действие:", reply_markup=_kb_main().as_markup())
        return
    await _send_chunked(msg, title, _lines_for_unvisited(items))
    await state.update_data(visit_num2id=[int(r["referral_id"]) for r in items])

    kb = InlineKeyboardBuilder()
    kb.button(text="✅ Отметить ВСЕ визиты", callback_data="adm:visit_all_ask")
    kb.button(text="❌ Отмена", callback_data="adm:visit_cancel")
    kb.adjust(2)
    await msg.answer("Или одним нажатием:", reply_markup=kb.as_markup())

    await msg.answer(
        "Введите номера для отметки визита (например: `3, 5-7, 12`).",
        parse_mode="Markdown"
    )
    await state.set_state(MarkVisited.waiting_numbers)


@router.callback_query(F.data == "adm:vis_all")
async def vis_all(cb: CallbackQuery, state: FSMContext):
    items = list_unvisited_all()
    await _start_mark_visit_flow(cb.message, items, state, "Не отмеченные визиты (все)")
    await cb.answer()


@router.callback_query(F.data == "adm:vis_choose_doctor")
async def vis_choose_doctor(cb: CallbackQuery):
    docs = list_doctors_with_counts()
    if not docs:
        await cb.message.answer("Список врачей пуст.")
        await cb.answer()
        return
    kb = InlineKeyboardBuilder()
    for i, d in enumerate(docs, 1):
        kb.button(text=f"{i}. {d['full_name']}", callback_data=f"adm:vis_doc:{d['doctor_id']}")
    kb.adjust(1)
    await cb.message.answer("Выберите врача:", reply_markup=kb.as_markup())
    await cb.answer()


@router.callback_query(F.data.startswith("adm:vis_doc:"))
async def vis_for_doctor(cb: CallbackQuery, state: FSMContext):
    doctor_id = int(cb.data.split(":")[-1])
    items = list_unvisited_by_doctor(doctor_id)
    await _start_mark_visit_flow(cb.message, items, state, "Не отмеченные визиты (по врачу)")
    await cb.answer()


@router.message(MarkVisited.waiting_numbers)
async def mark_visit_numbers(msg: Message, state: FSMContext):
    text = (msg.text or "").strip().lower()
    if text in {"отмена", "cancel", "stop"}:
        await state.clear()
        await msg.answer("Отменено.")
        await msg.answer("Выберите действие:", reply_markup=_kb_main().as_markup())
        return

    data = await state.get_data()
    num2id: list[int] = data.get("visit_num2id") or []
    if not num2id:
        await state.clear()
        await msg.answer("Истёк контекст. Откройте список заново.")
        await msg.answer("Выберите действие:", reply_markup=_kb_main().as_markup())
        return

    nums = _expand_numbers(text, len(num2id))
    if not nums:
        await msg.answer("Не распознал номера. Введите, например: `3, 5-7, 12`.", parse_mode="Markdown")
        return

    ref_ids = [num2id[i - 1] for i in nums]
    await state.update_data(visit_selected_ids=ref_ids, visit_selected_nums=nums)

    preview = ", ".join(map(str, nums))
    kb = InlineKeyboardBuilder()
    kb.button(text="✅ Отметить визит", callback_data="adm:visit_confirm")
    kb.button(text="❌ Отмена", callback_data="adm:visit_cancel")
    kb.adjust(2)

    await msg.answer(f"Подтвердите отметку визита для номеров [{preview}] (всего {len(nums)}).",
                     reply_markup=kb.as_markup())


@router.callback_query(F.data == "adm:visit_all_ask")
async def visit_all_ask(cb: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    ids = data.get("visit_num2id") or []
    if not ids:
        await cb.message.answer("Список пуст.")
        await cb.answer()
        return
    kb = InlineKeyboardBuilder()
    kb.button(text=f"✅ Подтвердить: {len(ids)}", callback_data="adm:visit_all_confirm")
    kb.button(text="❌ Отмена", callback_data="adm:visit_cancel")
    kb.adjust(2)
    await cb.message.answer("Отметить визит всем из показанного списка?", reply_markup=kb.as_markup())
    await cb.answer()


@router.callback_query(F.data == "adm:visit_all_confirm")
async def visit_all_confirm(cb: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    ids = data.get("visit_num2id") or []
    updated = mark_referrals_visited(ids)

    details = get_referrals_details(ids)
    lines = []
    for r in details:
        fio = r["patient_full_name"]
        bd = _fmt_birth_ru(r.get("patient_birth_date"))
        dadd = _fmt_ru_date(r.get("created_at"))
        lines.append(f"• {fio} {bd} — дата направления {dadd}")
    if lines:
        await cb.message.answer("Отмечены визиты:\n" + "\n".join(lines))

    await state.clear()
    await cb.message.answer(f"Отмечено визитов: {updated}.")
    await cb.message.answer("Выберите действие:", reply_markup=_kb_main().as_markup())
    await cb.answer()


@router.callback_query(F.data == "adm:visit_confirm")
async def visit_confirm(cb: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    ref_ids: list[int] = data.get("visit_selected_ids") or []
    updated = mark_referrals_visited(ref_ids)

    details = get_referrals_details(ref_ids)
    lines = []
    for r in details:
        fio = r["patient_full_name"]
        bd = _fmt_birth_ru(r.get("patient_birth_date"))
        dadd = _fmt_ru_date(r.get("created_at"))
        lines.append(f"• {fio} {bd} — дата направления {dadd}")
    if lines:
        await cb.message.answer("Отмечены визиты:\n" + "\n".join(lines))

    await state.clear()
    await cb.message.answer(f"Отмечено визитов: {updated}.")
    await cb.message.answer("Выберите действие:", reply_markup=_kb_main().as_markup())
    await cb.answer()


@router.callback_query(F.data == "adm:visit_cancel")
async def visit_cancel(cb: CallbackQuery, state: FSMContext):
    await state.clear()
    await cb.message.answer("Отменено.")
    await cb.message.answer("Выберите действие:", reply_markup=_kb_main().as_markup())
    await cb.answer()


# -------- Расчёт за визит --------

@router.callback_query(F.data == "adm:settle_menu")
async def settle_menu(cb: CallbackQuery, state: FSMContext):
    kb = InlineKeyboardBuilder()
    kb.button(text="Все, кто пришёл на визит", callback_data="adm:uns_all_ready")
    kb.button(text="По врачу", callback_data="adm:settle_choose_doctor")
    kb.button(text="⬅️ Назад", callback_data="adm:back")
    kb.adjust(1)
    await cb.message.answer("Выберите режим расчёта:", reply_markup=kb.as_markup())
    await cb.answer()


@router.callback_query(F.data == "adm:uns_all_ready")
async def ready_all(cb: CallbackQuery, state: FSMContext):
    items = list_ready_to_settle_all()
    await _start_mark_settle_flow(cb.message, items, state, "Расчёт за визит (все отметившиеся)")
    await cb.answer()


@router.callback_query(F.data == "adm:settle_choose_doctor")
async def settle_choose_doctor(cb: CallbackQuery, state: FSMContext):
    docs = list_doctors_with_counts()
    if not docs:
        await cb.message.answer("Список врачей пуст.")
        await cb.answer()
        return
    kb = InlineKeyboardBuilder()
    for i, d in enumerate(docs, 1):
        kb.button(text=f"{i}. {d['full_name']}", callback_data=f"adm:uns_doc_ready:{d['doctor_id']}")
    kb.adjust(1)
    await cb.message.answer("Выберите врача:", reply_markup=kb.as_markup())
    await cb.answer()


@router.callback_query(F.data.startswith("adm:uns_doc_ready:"))
async def ready_for_doctor(cb: CallbackQuery, state: FSMContext):
    doctor_id = int(cb.data.split(":")[-1])
    items = list_ready_to_settle_by_doctor(doctor_id)
    await _start_mark_settle_flow(cb.message, items, state, "Расчёт за визит (по врачу)")
    await cb.answer()


@router.message(MarkSettled.waiting_numbers)
async def settle_numbers(msg: Message, state: FSMContext):
    text = (msg.text or "").strip().lower()
    if text in {"отмена", "cancel", "stop"}:
        await state.clear()
        await msg.answer("Отменено.")
        await msg.answer("Выберите действие:", reply_markup=_kb_main().as_markup())
        return

    data = await state.get_data()
    num2id: list[int] = data.get("mark_num2id") or []
    if not num2id:
        await state.clear()
        await msg.answer("Истёк контекст. Откройте список заново.")
        await msg.answer("Выберите действие:", reply_markup=_kb_main().as_markup())
        return

    nums = _expand_numbers(text, len(num2id))
    if not nums:
        await msg.answer("Не распознал номера. Введите, например: `3, 5-7, 12`.", parse_mode="Markdown")
        return

    ref_ids = [num2id[i - 1] for i in nums]
    await state.update_data(mark_selected_ids=ref_ids, mark_selected_nums=nums)

    preview = ", ".join(map(str, nums))
    kb = InlineKeyboardBuilder()
    kb.button(text="✅ Выполнить расчёт", callback_data="adm:mark_confirm")
    kb.button(text="❌ Отмена", callback_data="adm:mark_cancel")
    kb.adjust(2)

    await msg.answer(f"Подтвердите расчёт для номеров [{preview}] (всего {len(nums)}).",
                     reply_markup=kb.as_markup())


@router.callback_query(F.data == "adm:settle_all_ask")
async def settle_all_ask(cb: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    ids = data.get("mark_num2id") or []
    if not ids:
        await cb.message.answer("Список пуст.")
        await cb.answer()
        return
    kb = InlineKeyboardBuilder()
    kb.button(text=f"✅ Подтвердить расчёт: {len(ids)}", callback_data="adm:settle_all_confirm")
    kb.button(text="❌ Отмена", callback_data="adm:mark_cancel")
    kb.adjust(2)
    await cb.message.answer("Рассчитать всех из показанного списка?", reply_markup=kb.as_markup())
    await cb.answer()


@router.callback_query(F.data == "adm:settle_all_confirm")
async def settle_all_confirm(cb: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    ids = data.get("mark_num2id") or []

    updated = settle_referrals(ids)

    details = get_referrals_details(ids)
    per_doctor: dict[int, list[dict]] = defaultdict(list)
    for r in details:
        per_doctor[int(r["doctor_tg_user_id"])].append(r)

    global doctor_notify_bot
    if doctor_notify_bot is not None:
        for tg_uid, rows in per_doctor.items():
            if not tg_uid:
                continue
            lines = []
            for r in rows:
                fio = r["patient_full_name"]
                bd = _fmt_birth_ru(r.get("patient_birth_date"))
                dadd = _fmt_ru_date(r.get("created_at"))
                lines.append(f"• {fio} {bd} — дата направления {dadd}")
            text = "✅ Рассчитались за:\n" + "\n".join(lines)
            try:
                await doctor_notify_bot.send_message(chat_id=tg_uid, text=text)
            except Exception as e:
                log.warning(f"Notify doctor {tg_uid} failed: {e}")

    await state.clear()
    await cb.message.answer(f"Рассчитано: {updated}.")
    await cb.message.answer("Выберите действие:", reply_markup=_kb_main().as_markup())
    await cb.answer()


@router.callback_query(F.data == "adm:mark_confirm")
async def mark_confirm(cb: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    ref_ids: list[int] = data.get("mark_selected_ids") or []

    updated = settle_referrals(ref_ids)
    await state.clear()

    details = get_referrals_details(ref_ids)
    per_doctor: dict[int, list[dict]] = defaultdict(list)
    for r in details:
        per_doctor[int(r["doctor_tg_user_id"])].append(r)

    global doctor_notify_bot
    if doctor_notify_bot is not None:
        for tg_uid, rows in per_doctor.items():
            if not tg_uid:
                continue
            lines = []
            for r in rows:
                fio = r["patient_full_name"]
                bd = _fmt_birth_ru(r.get("patient_birth_date"))
                dadd = _fmt_ru_date(r.get("created_at"))
                lines.append(f"• {fio} {bd} — дата направления {dadd}")
            text = "✅ Рассчитались за:\n" + "\n".join(lines)
            try:
                await doctor_notify_bot.send_message(chat_id=tg_uid, text=text)
            except Exception as e:
                log.warning(f"Notify doctor {tg_uid} failed: {e}")

    await cb.message.answer(f"Рассчитано: {updated}.")
    await cb.message.answer("Выберите действие:", reply_markup=_kb_main().as_markup())
    await cb.answer()


@router.callback_query(F.data == "adm:mark_cancel")
async def mark_cancel(cb: CallbackQuery, state: FSMContext):
    await state.clear()
    await cb.message.answer("Отменено.")
    await cb.message.answer("Выберите действие:", reply_markup=_kb_main().as_markup())
    await cb.answer()


# -------- Рассчитанные (текущий месяц) --------

@router.callback_query(F.data == "adm:settled_current")
async def settled_current(cb: CallbackQuery):
    rows = export_all_referrals()
    today = date.today()
    yy, mm = f"{today.year:04d}", f"{today.month:02d}"
    rows = [r for r in rows if str(r.get("settled")) == "1"
            and (r.get("settled_at") and r["settled_at"][:4] == yy and r["settled_at"][5:7] == mm)]
    if not rows:
        await cb.message.answer("За текущий месяц рассчитанных нет.")
        await cb.message.answer("Выберите действие:", reply_markup=_kb_main().as_markup())
        await cb.answer()
        return

    lines = []
    for i, r in enumerate(rows, 1):
        fio = r["patient_full_name"]
        bd = _fmt_birth_ru(r.get("patient_birth_date"))
        dadd = _fmt_ru_date(r.get("created_at"))
        dset = _fmt_ru_date(r.get("settled_at"))
        doc = r["doctor_full_name"]
        lines.append(f"{i}. {fio} {bd} — {doc} — направление: {dadd} — рассчитан: {dset}")
    await _send_chunked(cb.message, "Рассчитанные за текущий месяц:", lines)
    await cb.message.answer("Выберите действие:", reply_markup=_kb_main().as_markup())
    await cb.answer()


# -------- Экспорт (Excel) --------

@router.callback_query(F.data == "adm:export")
async def export_xlsx(cb: CallbackQuery):
    doctors = export_doctors_overview()
    refs = export_all_referrals()

    wb = Workbook()
    ws = wb.active
    ws.title = "Врачи"
    ws.append(["#", "Врач", "Всего", "Не рассчитано", "Рассчитано"])
    for i, d in enumerate(doctors, 1):
        ws.append([i, d["doctor_full_name"], int(d["total"] or 0), int(d["unsettled"] or 0), int(d["settled"] or 0)])
    _autosize(ws)

    ws2 = wb.create_sheet("Все направления")
    ws2.append([
        "#", "Доктор", "Пациент", "Дата рождения",
        "Дата направления", "Статус", "Дата расчёта",
        "Визит", "Дата визита"
    ])
    for i, r in enumerate(refs, 1):
        status = "рассчитан" if int(r["settled"] or 0) == 1 else "не рассчитан"
        visited_flag = "отмечен" if int(r["visited"] or 0) == 1 else "не отмечен"
        ws2.append([
            i,
            r["doctor_full_name"],
            r["patient_full_name"],
            _fmt_birth_ru(r.get("patient_birth_date")),
            _fmt_ru_date(r.get("created_at")),
            status,
            _fmt_ru_date(r.get("settled_at")),
            visited_flag,
            _fmt_ru_date(r.get("visited_at")),
        ])
    _autosize(ws2)

    ws3 = wb.create_sheet("Рейтинг врачей")
    top_ref = top_doctors_by_referrals(limit=100)
    ws3.append(["ТОП по направлениям (весь период)"])
    ws3.append(["#", "Врач", "Направлений"])
    for i, r in enumerate(top_ref, 1):
        ws3.append([i, r["doctor_full_name"], int(r["cnt"] or 0)])
    ws3.append([])
    top_vis = top_doctors_by_visits(limit=100)
    ws3.append(["ТОП по визитам (весь период)"])
    ws3.append(["#", "Врач", "Визитов"])
    for i, r in enumerate(top_vis, 1):
        ws3.append([i, r["doctor_full_name"], int(r["cnt"] or 0)])
    _autosize(ws3)

    exports_dir = ROOT / "data" / "exports"
    exports_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    fname = exports_dir / f"clinic_export_{ts}.xlsx"
    wb.save(fname)

    await cb.message.answer_document(
        document=FSInputFile(str(fname)),
        caption=f"Экспорт от {date.today().strftime('%d.%m.%Y')}"
    )
    await cb.answer()


# -------- Удалить пациента(ов) --------

@router.callback_query(F.data == "adm:del_patients")
async def del_patients_start(cb: CallbackQuery, state: FSMContext):
    data = list_patients_with_ref_counts()
    if not data:
        await cb.message.answer("Пациентов нет.")
        await cb.answer()
        return

    lines = []
    for i, p in enumerate(data, 1):
        lines.append(f"{i}. {p['full_name']} — {_fmt_birth_ru(p.get('birth_date'))} — направлений: {int(p['referrals_count'] or 0)}")
    await _send_chunked(cb.message, "Пациенты (для удаления по номерам):", lines)

    num2id = [int(p["patient_id"]) for p in data]
    await state.update_data(delp_num2id=num2id, delp_lines=lines)

    await cb.message.answer(
        "Введите номера пациентов для удаления (например: `3, 5-7, 12`).\n"
        "Внимание: их направления будут удалены каскадно.",
        parse_mode="Markdown",
        reply_markup=_kb_back_menu().as_markup()
    )
    await state.set_state(DelPatients.waiting_numbers)
    await cb.answer()


@router.message(DelPatients.waiting_numbers)
async def del_patients_numbers(msg: Message, state: FSMContext):
    text = (msg.text or "").strip().lower()
    if text in {"отмена", "cancel", "stop"}:
        await state.clear()
        await msg.answer("Отменено.")
        await msg.answer("Выберите действие:", reply_markup=_kb_main().as_markup())
        return

    data = await state.get_data()
    num2id: list[int] = data.get("delp_num2id") or []
    lines: list[str] = data.get("delp_lines") or []
    if not num2id:
        await state.clear()
        await msg.answer("Истёк контекст. Откройте список заново.")
        await msg.answer("Выберите действие:", reply_markup=_kb_main().as_markup())
        return

    nums = _expand_numbers(text, len(num2id))
    if not nums:
        await msg.answer("Не распознал номера. Введите, например: `3, 5-7, 12`.", parse_mode="Markdown")
        return

    sel_ids = [num2id[i - 1] for i in nums]
    preview_lines = [lines[i - 1] for i in nums][:20]
    await state.update_data(delp_selected_ids=sel_ids, delp_selected_nums=nums)

    kb = InlineKeyboardBuilder()
    kb.button(text=f"✅ Удалить: {len(sel_ids)}", callback_data="adm:del_patients_confirm")
    kb.button(text="❌ Отмена", callback_data="adm:del_patients_cancel")
    kb.adjust(2)

    await msg.answer("К удалению:\n" + "\n".join(preview_lines), disable_web_page_preview=True)
    await msg.answer("Подтвердите удаление выбранных пациентов.", reply_markup=kb.as_markup())


@router.callback_query(F.data == "adm:del_patients_confirm")
async def del_patients_confirm(cb: CallbackQuery, state: FSMContext):
    data = await state.get_data()
    ids: list[int] = data.get("delp_selected_ids") or []
    deleted = delete_patients(ids)
    await state.clear()
    await cb.message.answer(f"Удалено пациентов: {deleted}. (Связанные направления удалены каскадно)")
    await cb.message.answer("Выберите действие:", reply_markup=_kb_main().as_markup())
    await cb.answer()


@router.callback_query(F.data == "adm:del_patients_cancel")
async def del_patients_cancel(cb: CallbackQuery, state: FSMContext):
    await state.clear()
    await cb.message.answer("Отменено.")
    await cb.message.answer("Выберите действие:", reply_markup=_kb_main().as_markup())
    await cb.answer()


# -------- Back --------

@router.callback_query(F.data == "adm:back")
async def go_back(cb: CallbackQuery):
    await cb.message.answer("Выберите действие:", reply_markup=_kb_main().as_markup())
    await cb.answer()


# -------- Entrypoint --------

async def main():
    migrate()
    admin_bot = Bot(get_admin_token())

    global doctor_notify_bot
    doctor_notify_bot = Bot(get_doctor_token())

    dp = Dispatcher(storage=MemoryStorage())
    dp.include_router(router)

    log.info("Admin bot starting...")
    await dp.start_polling(admin_bot)


if __name__ == "__main__":
    try:
        asyncio.run(main())
    except Exception:
        log.exception("Fatal error in admin bot")
        raise
